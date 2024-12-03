import tkinter as tk
import csv
import openpyxl
from PIL import Image, ImageTk
from collections import defaultdict


class QuestionnaireApp:
    def __init__(self, window, background_path, csv_file):
        self.window = window
        self.window.geometry("600x400")
        self.window.resizable(False, False)

        # Hintergrundbild laden
        self.bg_image = Image.open(background_path)
        self.bg_image = self.bg_image.resize((600, 400), Image.Resampling.LANCZOS)
        self.bg_photo = ImageTk.PhotoImage(self.bg_image)

        self.bg_label = tk.Label(window, image=self.bg_photo)
        self.bg_label.place(x=0, y=0, relwidth=1, relheight=1)

        # Liste von Antworten
        self.answers = []

        # CSV-Datei laden
        self.questions = self.load_csv(csv_file)

        # Variable zum Speichern der Antwort
        self.user_answer = tk.StringVar()

        self.current_question = 0
        self.total_questions = len(self.questions)

        # Start der ersten Frage
        self.show_question()

    def load_csv(self, csv_file):
        questions = []
        with open(csv_file, newline='', encoding='utf-8') as csvfile:
            reader = csv.reader(csvfile, delimiter=';')
            for row in reader:
                if row:
                    questions.append({
                        'id': row[0],
                        'type': int(row[1]),
                        'options': row[2].split(','),
                        'question': row[3],
                        'tag': row[4],
                        'condition': row[5]
                    })
        return questions

    def create_button(self, text, command):
        return tk.Button(self.window, text=text, command=command, width=32, height=2, bg="#3B3D51", fg="#ffffff",
                         font=("Arial", 12, "bold"), activebackground="#7579A0", activeforeground="#ffffff", relief=tk.RAISED,
                         bd=3, cursor="hand2", highlightthickness=0)

    def create_entry(self):
        return tk.Entry(self.window, font=("Arial", 12, "bold"), width=36, bg="#3B3D51", fg="#ffffff", insertbackground="#ffffff", bd=3, relief=tk.RAISED)

    def show_question(self):
        question = self.questions[self.current_question]
        self.window.title(f"Frage {self.current_question + 1}/{self.total_questions}")

        # Lösche vorherige Widgets (außer Hintergrund)
        for widget in self.window.winfo_children():
            if widget != self.bg_label:
                widget.destroy()

        # Frage anzeigen
        question_label = tk.Label(self.window, text=question['question'], font=("Arial", 16, "bold"), fg="#ffffff", bg="#3B3D51")
        question_label.pack(pady=(100, 10), padx=10)

        # Funktionsaufruf für den Button oder Eingabefeld
        if question['type'] == 0:
            self.create_yes_no_buttons()
        elif question['type'] == 1:
            self.create_option_buttons(question['options'])
        elif question['type'] == 99:
            self.create_input_field()

    def create_yes_no_buttons(self):
        button_yes = self.create_button("Ja", lambda: self.record_answer("Ja"))
        button_no = self.create_button("Nein", lambda: self.record_answer("Nein"))
        button_yes.pack(pady=5)
        button_no.pack(pady=5)

    def create_option_buttons(self, options):
        for option in options:
            button = self.create_button(option, lambda opt=option: self.record_answer(opt))
            button.pack(pady=5)

    def create_input_field(self):
        entry = self.create_entry()
        entry.pack(pady=5)
        submit_button = self.create_button("Absenden", lambda: self.record_answer(entry.get()))
        submit_button.pack(pady=5)

    def record_answer(self, answer):
        question = self.questions[self.current_question]
        # Antwort und zugehörige Daten speichern
        self.answers.append({
            'id': question['id'],
            'answer': answer,
            'tag': question['tag'],
            'condition': question['condition'],
        })

        # Antwort speichern und nächste Frage aufrufen
        self.current_question += 1
        if self.current_question < self.total_questions:
            self.show_question()
        else:
            self.show_finish_screen()

    def show_finish_screen(self):
        for widget in self.window.winfo_children():
            if widget != self.bg_label:
                widget.destroy()

        # Scrollbarer Canvas
        canvas = tk.Canvas(self.window, bg="#3B3D51", highlightthickness=0)
        scrollbar = tk.Scrollbar(self.window, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas, bg="#3B3D51")

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="x", expand=True)
        scrollbar.pack(side="right",fill="y")

        finish_label = tk.Label(scrollable_frame, text="Vielen Dank für Ihre Teilnahme!", font=("Arial", 20, "bold"), fg="#ffffff", bg="#3B3D51")
        finish_label.pack(pady=20)

        # Gruppiere die Antworten nach Tag und zeige Ergebnisse
        grouped_answers = self.group_answers_by_tag()

        for tag, group in grouped_answers.items():
            group_label = tk.Label(scrollable_frame, text=f"Gruppe: {tag}", font=("Arial", 16, "bold"), fg="#ffffff", bg="#3B3D51")
            group_label.pack(pady=10)

            for answer in group:
                condition_result = self.evaluate_condition(answer['condition'])
                answer_label = tk.Label(scrollable_frame, text=f"Antwort: {answer['answer']} - Bedingung erfüllt: {condition_result}",
                                        font=("Arial", 14), fg="#ffffff", bg="#3B3D51")
                answer_label.pack(pady=5)

        # Button zum Beenden
        #exit_button = self.create_button("Beenden", self.window.quit)
        #exit_button.pack(pady=30)

        # Speichere die Ergebnisse in einer Excel-Datei
        self.save_to_excel(grouped_answers)

    def save_to_excel(self, grouped_answers):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Gruppe", "Antwort", "Bedingung", "Bedingung erfüllt"])

        for tag, group in grouped_answers.items():
            for answer in group:
                condition_result = self.evaluate_condition(answer['condition'])
                ws.append([answer['id'], tag, answer['answer'], answer['condition'], condition_result])

        file_name = "exel/Fragen_Antworten.xlsx"
        wb.save(file_name)
        print(f"Ergebnisse wurden in {file_name} gespeichert.")

    def group_answers_by_tag(self):
        # Gruppiere die Antworten nach "Tag"
        grouped_answers = defaultdict(list)
        for answer in self.answers:
            grouped_answers[answer['tag']].append(answer)
        return grouped_answers

    def evaluate_condition(self, condition):
        # Map placeholders (e.g., 'a', 'b', 'c') to actual answers based on their index
        placeholder_map = {}
        for idx, answer in enumerate(self.answers):
            placeholder_map[chr(97 + idx)] = answer['answer']  # 'a', 'b', 'c', ...

        try:
            # Replace placeholders in the condition
            for placeholder, value in placeholder_map.items():
                condition = condition.replace(placeholder, f'"{value}"')

            # Evaluate the condition
            return eval(condition)
        except Exception as e:
            return f"Fehler ({e})"

    def run(self):
        self.window.mainloop()


def main():
    # Erstelle die GUI-Anwendung
    window = tk.Tk()

    # Erstelle die App
    app = QuestionnaireApp(window, background_path="background.png", csv_file="exel/fragen.csv")
    app.run()


if __name__ == "__main__":
    main()
